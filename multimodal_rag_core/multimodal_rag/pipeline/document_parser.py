from __future__ import annotations

import logging
import re
from pathlib import Path

from multimodal_rag.config import Settings
from multimodal_rag.pipeline.chunking import Chunk, text_to_chunks

logger = logging.getLogger(__name__)

DOCLING_SUFFIXES = {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".md"}

# Warn-once flag: a missing Docling install otherwise degrades silently
# (per-file debug logs are invisible at the default INFO level).
_DOCLING_MISSING_WARNED = False


def parse_document(path: Path, settings: Settings) -> list[Chunk]:
    text = _extract_text(path, settings)
    if not text or not text.strip():
        return []
    return text_to_chunks(text, settings.DOC_CHUNK_SIZE, settings.DOC_CHUNK_OVERLAP)


def _extract_text(path: Path, settings: Settings) -> str:
    if settings.DOC_USE_OCR and path.suffix.lower() in DOCLING_SUFFIXES:
        markdown = _docling_markdown(path)
        if markdown:
            return markdown
    return _fallback_text(path)


def _docling_markdown(path: Path) -> str:
    global _DOCLING_MISSING_WARNED
    try:
        from docling.document_converter import DocumentConverter
    except Exception:
        if not _DOCLING_MISSING_WARNED:
            _DOCLING_MISSING_WARNED = True
            logger.warning(
                "Docling is not installed, so DOC_USE_OCR has no effect: "
                "scanned/image-only PDFs will extract no text. "
                "Install it with `pip install docling` (pulls torch + OCR models)."
            )
        return ""
    try:
        result = DocumentConverter().convert(str(path))
        return result.document.export_to_markdown()
    except Exception as exc:
        logger.warning("Docling failed for %s: %s", path, exc)
        return ""


def _fallback_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv", ".json", ".log", ".rst", ".xml"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix in {".html", ".htm"}:
            raw = path.read_text(encoding="utf-8", errors="ignore")
            return re.sub(r"<[^>]+>", " ", raw)
        if suffix == ".pdf":
            return _pdf_text(path)
        if suffix == ".docx":
            import docx

            document = docx.Document(str(path))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        if suffix in {".xlsx", ".xls"}:
            return _spreadsheet_text(path)
    except Exception as exc:
        logger.warning("Fallback parse failed for %s: %s", path, exc)
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""


def _pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception:
        from PyPDF2 import PdfReader  # type: ignore
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _spreadsheet_text(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            lines.append(" | ".join("" if cell is None else str(cell) for cell in row))
    return "\n".join(lines)
